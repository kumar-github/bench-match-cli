"""
Standalone, thin companion to write_shortlist_workbook.py.

Produces a narrow, one-row-per-employee view for quick scanning:
  - the single best Strong match per employee, or if none exists, the
    single best Good match
  - employees with neither a Strong nor a Good match are left out of this
    file entirely (they're still fully visible in the main workbook)
  - raw literal skill-cluster text (not the persona label) so a plain
    text search for a specific skill (e.g. "Spring Boot") actually finds
    something

This does NOT touch engine.py or build_matches.py, and does not recompute
any score or tier -- it only reads data/match_rows.csv (already produced
by build_matches.py) plus the two classified CSVs, picks one row per
employee, and renders it. Run it any time after run.py (or the individual
pipeline steps) has produced data/match_rows.csv.

Usage:
    python generate-thin.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "output"

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=10)
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="595959")
STRONG_FILL = PatternFill("solid", fgColor="C6EFCE")
GOOD_FILL = PatternFill("solid", fgColor="FFEB9C")
STRONG_FONT = Font(name=FONT_NAME, size=10, bold=True, color="006100")
GOOD_FONT = Font(name=FONT_NAME, size=10, bold=True, color="9C6500")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SCORE_FRACTION = {"Strong": 1.0, "Good": 0.75}  # real numbers, displayed via a percent format

COLUMNS = [
    ("Demand ID", 13),
    ("Demand Cluster Name", 42),
    ("Demand Location", 15),
    ("Demand Sub Band", 15),
    ("Employee ID", 13),
    ("Employee Name", 26),
    ("Employee Location", 15),
    ("Employee Cluster", 42),
    ("Employee Sub Band", 15),
    ("Matching Score", 13),
    ("Reasoning (one-line)", 60),
    ("MAS Mapping Check", 30),
    ("Sub-Capability Check", 30),
    ("Additional Request (manual review only)", 50),
]


def pick_best_row(emp_rows):
    """Best Strong if any exist, else best Good, else None.
    Tie-break matches the existing Good-tier ordering in build_matches.py:
    most urgent ageing bucket first, then more balance positions."""
    strong = emp_rows[emp_rows["overall_tier"] == "Strong"]
    pool = strong if len(strong) else emp_rows[emp_rows["overall_tier"] == "Good"]
    if len(pool) == 0:
        return None
    pool = pool.sort_values(by=["ageing_rank", "balance_positions"], ascending=[True, False])
    return pool.iloc[0]


def build():
    match_path = DATA_DIR / "match_rows.csv"
    if not match_path.exists():
        raise FileNotFoundError(
            f"{match_path} not found. Run run.py (or classify_inputs.py + "
            f"build_matches.py) first -- this script only reformats their output."
        )

    df = pd.read_csv(match_path)
    supply = pd.read_csv(DATA_DIR / "phase1_supply_classified.csv")
    demand = pd.read_csv(DATA_DIR / "demand_classified.csv")

    # Raw literal skill-cluster text, keyed for the join back onto match_rows.
    emp_cluster = supply.set_index("Employee Code")["Skill Cluster"]
    dem_cluster = demand.set_index("Job Requisition ID")["Final Skill Cluster"]

    rows = []
    dropped_no_strong_good = 0
    for emp_code, emp_rows in df.groupby("employee_code"):
        best = pick_best_row(emp_rows)
        if best is None:
            dropped_no_strong_good += 1
            continue
        demand_id = best["job_req_id"]
        try:
            demand_id = int(demand_id)  # store as a real number, not text
        except (TypeError, ValueError):
            pass  # leave non-numeric IDs (if any) as-is
        rows.append(dict(
            demand_id=demand_id,
            demand_cluster=dem_cluster.get(best["job_req_id"], "(not found)"),
            demand_location=best["demand_location"],
            demand_sub_band=best["demand_band"],
            employee_id=int(emp_code) if pd.notna(emp_code) else emp_code,
            employee_name=best["employee_name"],
            employee_location=best["employee_location"],
            employee_cluster=emp_cluster.get(emp_code, "(not found)"),
            employee_sub_band=best["employee_band"],
            tier=best["overall_tier"],
            reasoning=best["one_liner"],
            mas_mapping_note=best.get("mas_mapping_note") if pd.notna(best.get("mas_mapping_note")) else "",
            sub_capability_note=best.get("employee_sub_capability_note") if pd.notna(best.get("employee_sub_capability_note")) else "",
            additional_request=best.get("additional_request") if pd.notna(best.get("additional_request")) else "",
        ))

    out_df = pd.DataFrame(rows)
    print(f"Employees with a Strong or Good match: {len(out_df)}")
    print(f"Employees dropped (no Strong/Good at all): {dropped_no_strong_good}")
    if len(out_df):
        print(out_df["tier"].value_counts())

    write_workbook(out_df)


def write_workbook(df):
    wb = Workbook()
    ws = wb.active
    ws.title = "Thin Shortlist"

    ws["A1"] = "Phase 1 Matching Shortlist — Thin View"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("One row per employee: their single best Strong match, or if none, "
                "their single best Good match. Employees with neither are not listed here "
                "-- see the full workbook for those.")
    ws["A2"].font = SUBTITLE_FONT

    header_row = 4
    for c, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=c, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = f"A{header_row + 1}"

    row = header_row + 1
    for _, r in df.iterrows():
        tier_fill = STRONG_FILL if r["tier"] == "Strong" else GOOD_FILL
        tier_font = STRONG_FONT if r["tier"] == "Strong" else GOOD_FONT
        values = [
            r["demand_id"], r["demand_cluster"], r["demand_location"], r["demand_sub_band"],
            r["employee_id"], r["employee_name"], r["employee_location"], r["employee_cluster"],
            r["employee_sub_band"], SCORE_FRACTION[r["tier"]], r["reasoning"],
            r["mas_mapping_note"], r["sub_capability_note"], r["additional_request"],
        ]
        for c, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=val)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (11, 13, len(COLUMNS))))
            if c == 1:  # Demand ID -- a real number, right-aligned like any ID column
                cell.alignment = Alignment(horizontal="right", vertical="top")
            if c == 10:  # Matching Score -- real number (1.0 / 0.75), displayed as a percent
                cell.number_format = "0%"
                cell.fill = tier_fill
                cell.font = tier_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if c == 12 and val:  # MAS Mapping Check -- only ever has a value when flagged
                cell.font = Font(name=FONT_NAME, size=10, bold=True, color="9C6500")
        row += 1

    OUTPUT_DIR.mkdir(exist_ok=True)
    out_path = OUTPUT_DIR / "phase1-matching-shortlist-thin.xlsx"
    wb.save(str(out_path))
    print(f"Workbook written to {out_path}")


if __name__ == "__main__":
    build()
