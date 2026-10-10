"""
Generates demand-classification-review-worklist.xlsx from the Step 1 outputs
(data/demand_classified.csv, data/phase1_supply_classified.csv).

Written 2026-10-01 as a reusable, re-runnable script -- the first version of this workbook
(2026-09-30) was produced ad hoc during an analysis session and never saved as a script, which is
exactly the kind of thing this project's own README already warns against (see the "Lesson for
whoever picks this up next" note on the shortlist-workbook re-port bug). Run this any time after
`python run.py` (or just `classify_inputs.main()`) to regenerate the worklist against current data
-- e.g. after the 2026-10-01 MAS Mapping scope gate change, which is exactly why this run was
needed: the previous workbook's single highest-impact finding (the Tpm-tagged
"Java,Microservices,React.js" rows) is now moot, since those rows are excluded from classification
entirely rather than being an open judgment call.

Methodology (unchanged from the original 2026-09-30 pass -- see
demand-supply-mapping-requirements.md's "Review worklist" section for the full rationale):
  - Only classified, in-scope demand rows (persona not null) are considered -- rows now cleared by
    either the bad-band check or the MAS Mapping scope gate are excluded automatically, same as
    they're excluded from matching.
  - A row is "flagged" if mas_mapping_flag OR frontend_anchor_flag is True, OR its classification
    status falls in one of the three Weak-tier buckets (framework-unconfirmed, weak MERN, or
    dual-acceptable).
  - Reviewed at the level of distinct Final Skill Cluster phrases, not individual rows -- one
    judgment call per phrase.
  - Sorted by flagged open positions (Balance Positions -- i.e. positions still open, not the
    original ask or the "Actionable Position" sub-metric), highest first.
  - Rows making up the first 80% of cumulative flagged-position volume are highlighted amber.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "output"

WEAK_TIER_MARKERS = ("framework-unconfirmed", "weak MERN", "dual-acceptable")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
AMBER_FILL = PatternFill("solid", fgColor="FFE699")
TITLE_FONT = Font(bold=True, size=14)
WRAP = Alignment(wrap_text=True, vertical="top")


def _row_reason(row):
    """Builds the 'Why flagged' text for a single demand row -- mirrors the original
    2026-09-30 pass's reason text exactly (mas_mapping_note, then frontend_anchor_note, then the
    Weak-tier status text, each only if applicable, joined with ' | ')."""
    parts = []
    if row.get("mas_mapping_flag") and isinstance(row.get("mas_mapping_note"), str):
        parts.append(row["mas_mapping_note"])
    if row.get("frontend_anchor_flag") and isinstance(row.get("frontend_anchor_note"), str):
        parts.append(row["frontend_anchor_note"])
    status = row.get("status")
    if isinstance(status, str) and any(marker in status for marker in WEAK_TIER_MARKERS):
        parts.append(status)
    return " | ".join(parts)


def build_review_worklist(demand_classified: pd.DataFrame) -> pd.DataFrame:
    classified = demand_classified[demand_classified["persona"].notna()].copy()

    weak_tier = classified["status"].apply(
        lambda s: isinstance(s, str) and any(m in s for m in WEAK_TIER_MARKERS)
    )
    flagged = classified["mas_mapping_flag"].fillna(False) | classified["frontend_anchor_flag"].fillna(False) | weak_tier
    classified["_flagged"] = flagged
    classified["_reason"] = classified.apply(_row_reason, axis=1)
    classified["_positions"] = classified["Balance Positions"].fillna(0)

    rows = []
    for phrase, group in classified.groupby("Final Skill Cluster"):
        flagged_group = group[group["_flagged"]]
        if len(flagged_group) == 0:
            continue
        mas_values = sorted({
            (str(v).strip() if pd.notna(v) else "None") for v in group["MAS Mapping"]
        })
        reasons = [r for r in flagged_group["_reason"].unique() if r]
        reason_text = " ;; ".join(reasons)
        if len(reason_text) > 500:
            reason_text = reason_text[:500] + "..."
        rows.append({
            "Final Skill Cluster (demand phrase)": phrase,
            "Engine persona": group["persona"].iloc[0],
            "Flagged rows / total rows using this phrase": f"{len(flagged_group)} / {len(group)}",
            "Flagged open positions / total positions": (
                f"{int(flagged_group['_positions'].sum())} / {int(group['_positions'].sum())}"
            ),
            "MAS Mapping values seen on these rows": ", ".join(mas_values),
            "Why flagged": reason_text,
            "Reviewer decision (Confirm engine / Override to ___ / Needs more info)": None,
            "Reviewer notes": None,
            "_sort_positions": int(flagged_group["_positions"].sum()),
        })

    worklist = pd.DataFrame(rows).sort_values("_sort_positions", ascending=False)
    return worklist.drop(columns="_sort_positions"), classified


def _autosize_and_header(ws, ncols, widths):
    for col in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    ws.freeze_panes = "A2"


def write_workbook(worklist: pd.DataFrame, classified_demand: pd.DataFrame, supply_classified: pd.DataFrame):
    wb = Workbook()

    # --- Sheet 1: Review Worklist ---
    ws = wb.active
    ws.title = "Review Worklist"
    ws.append(["Demand Classification Review Worklist"])
    ws["A1"].font = TITLE_FONT
    ws.append([
        "Every distinct Final Skill Cluster phrase (in classified, in-scope demand) that the "
        "engine itself flagged as uncertain -- MAS Mapping disagrees with the derived persona, a "
        "frontend framework is named without its backend anchor confirmed, or the pairing landed "
        "in a Weak-tier bucket (framework-unconfirmed / weak MERN / dual-acceptable). Sorted by "
        "flagged open positions, highest first -- review from the top down for the best return on "
        "time spent. Rows highlighted amber are the phrases making up the first 80% of flagged "
        "position volume."
    ])
    ws.append([])
    header_row = 4
    ws.append(list(worklist.columns))
    total_flagged_positions = 0
    for _, r in worklist.iterrows():
        flagged_positions = int(r["Flagged open positions / total positions"].split(" / ")[0])
        total_flagged_positions += flagged_positions
    running = 0
    amber_rows = set()
    for idx, r in worklist.iterrows():
        flagged_positions = int(r["Flagged open positions / total positions"].split(" / ")[0])
        running += flagged_positions
        if total_flagged_positions and running <= 0.8 * total_flagged_positions:
            amber_rows.add(idx)
    for i, (idx, r) in enumerate(worklist.iterrows()):
        ws.append(list(r.values))
        if idx in amber_rows:
            for cell in ws[header_row + 1 + i]:
                cell.fill = AMBER_FILL
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = WRAP
    widths = [35, 16, 16, 16, 28, 60, 30, 20]
    for col in range(1, len(worklist.columns) + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    for cell in ws[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    ws.freeze_panes = f"A{header_row + 1}"

    # --- Sheet 2: Supply Sub-Capability flags ---
    ws2 = wb.create_sheet("Supply Sub-Capability flags")
    ws2.append(["Supply-side Sub-Capability review flags"])
    ws2["A1"].font = TITLE_FONT
    ws2.append([
        "Employees where the source file's own Sub-Capability tag disagrees with the persona the "
        "engine derived from Skill Cluster. Informational only -- never used to filter or score -- "
        "but a real place stale HR tagging shows up."
    ])
    ws2.append([])
    cols2 = ["Employee Code", "Employee Name", "Skill Cluster (raw)", "Engine persona", "Sub-Capability (raw)", "Sub Band"]
    ws2.append(cols2)
    flagged_supply = supply_classified[supply_classified["sub_capability_flag"] == True]  # noqa: E712
    for _, r in flagged_supply.iterrows():
        ws2.append([
            r["Employee Code"], r["Employee Name"], r["Skill Cluster"], r["persona"], r["Sub-Capability"], r["Sub Band"]
        ])
    widths2 = [16, 26, 36, 16, 20, 12]
    for col in range(1, len(cols2) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = widths2[col - 1]
    for cell in ws2[4]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws2.freeze_panes = "A5"

    # --- Sheet 3: Summary ---
    ws3 = wb.create_sheet("Summary")
    ws3.append(["Review Worklist -- Summary & Methodology"])
    ws3["A1"].font = TITLE_FONT
    n_classified = len(classified_demand)
    total_positions = int(classified_demand["_positions"].sum())
    n_flagged_rows = int(classified_demand["_flagged"].sum())
    flagged_positions_total = int(classified_demand.loc[classified_demand["_flagged"], "_positions"].sum())
    n_phrases_total = classified_demand["Final Skill Cluster"].nunique()
    n_phrases_flagged = len(worklist)
    pct_flagged_rows = round(100 * n_flagged_rows / n_classified) if n_classified else 0
    pct_flagged_positions = round(100 * flagged_positions_total / total_positions) if total_positions else 0

    top5 = worklist.head(5)
    top5_positions = sum(int(r["Flagged open positions / total positions"].split(" / ")[0]) for _, r in top5.iterrows())
    pct_top5 = round(100 * top5_positions / flagged_positions_total) if flagged_positions_total else 0

    lines = [
        None,
        f"Classified, in-scope demand rows: {n_classified}, totalling {total_positions} open positions "
        f"('open positions' = Balance Positions; excludes rows cleared by the bad-band check or the "
        f"2026-10-01 MAS Mapping scope gate -- those are no longer an open review question, they're "
        f"out of Phase 1 scope entirely).",
        f"Rows carrying at least one engine uncertainty flag: {n_flagged_rows} ({pct_flagged_rows}%), "
        f"totalling {flagged_positions_total} open positions ({pct_flagged_positions}% of all classified positions).",
        f"Collapsed to distinct Final Skill Cluster phrases: {n_phrases_flagged} of {n_phrases_total} "
        f"distinct phrases need review -- reviewing a phrase once covers every row that shares it.",
        None,
        "Concentration -- why sorting by position volume matters:",
        f"The top 5 flagged phrases alone account for {top5_positions} of {flagged_positions_total} "
        f"flagged positions ({pct_top5}%). Rows highlighted amber on the Review Worklist tab are the "
        f"phrases making up the first 80% of flagged position volume -- reviewing just those closes "
        f"most of the gap.",
        None,
        "What changed since the 2026-09-30 version of this workbook:",
        "The single highest-impact finding in that version -- 'Java,Microservices,React.js', 32 of "
        "whose 79 rows were tagged MAS Mapping='Tpm' (Technical Project Management) yet classified "
        "Fullstack Java purely from skill text -- is now resolved, not an open review item. The "
        "2026-10-01 MAS Mapping scope gate excludes Tpm (and Testing/Polyglot/Devops & Sre/"
        "Integration/Mainframe/Architects & Emerging) from Phase 1 entirely; those rows no longer "
        "appear anywhere in this workbook because they're not classified at all -- see "
        "demand-supply-mapping-requirements.md's 'MAS Mapping scoping' section for the full decision.",
        None,
        "How to use the Review Worklist tab:",
        "One row per distinct demand phrase, not per requisition -- a decision here applies to "
        "every row sharing that phrase. Fill in 'Reviewer decision' (Confirm the engine's persona / "
        "Override to a specific persona / Needs more info from the demand team) and any notes. "
        "Decisions can be fed back into engine.py's classification rules as new, permanent "
        "exceptions -- the same discipline already used to resolve every prior sanity-check "
        "finding in this project.",
    ]
    for line in lines:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 120
    for row in ws3.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    return wb


def main():
    demand_classified = pd.read_csv(DATA_DIR / "demand_classified.csv")
    supply_classified = pd.read_csv(DATA_DIR / "phase1_supply_classified.csv")

    worklist, classified_demand = build_review_worklist(demand_classified)
    wb = write_workbook(worklist, classified_demand, supply_classified)

    OUTPUT_DIR.mkdir(exist_ok=True)
    out_path = OUTPUT_DIR / "demand-classification-review-worklist.xlsx"
    wb.save(out_path)

    n_classified = len(classified_demand)
    n_flagged = int(classified_demand["_flagged"].sum())
    print(f"Classified in-scope demand rows: {n_classified}")
    print(f"Flagged rows: {n_flagged} across {len(worklist)} distinct phrases")
    print(f"Workbook written to {out_path}")


if __name__ == "__main__":
    main()
